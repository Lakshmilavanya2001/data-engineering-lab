"""Build the Lab 18 overview deck on the HPE/NVIDIA Light template with python-pptx.
Usage: python build_ppt.py <template.pptx> <out.pptx>"""
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TPL, OUT = sys.argv[1], sys.argv[2]
ARCH = '/home/administrator/pcai-dataeng-lab/guide/architecture/lab18_architecture.png'
GREEN=RGBColor(0x01,0xA9,0x82); INK=RGBColor(0x29,0x2D,0x3A); GREY=RGBColor(0x53,0x5C,0x66)
LIGHT=RGBColor(0xB1,0xB9,0xBE); PANEL=RGBColor(0xF2,0xF4,0xF5); PURPLE=RGBColor(0x77,0x64,0xFC)
BLUE=RGBColor(0x00,0x70,0xF8); WHITE=RGBColor(0xFF,0xFF,0xFF); ORANGE=RGBColor(0xC6,0x51,0x1C)
FOOT='Confidential | For training purposes only'

prs=Presentation(TPL)
# drop the template's sample slides, keep masters/layouts
sldIdLst=prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
L={l.name:l for l in prs.slide_masters[0].slide_layouts}

def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx==idx: return p
    return None
def set_text(shape, text, size=None, bold=None, color=None):
    tf=shape.text_frame; tf.clear(); p=tf.paragraphs[0]; r=p.add_run(); r.text=text
    if size: r.font.size=Pt(size)
    if bold is not None: r.font.bold=bold
    if color is not None: r.font.color.rgb=color
    return r
def tb(slide, x,y,w,h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.05):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=s.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(0.03); tf.vertical_anchor=anchor
    lines=text if isinstance(text,list) else [text]
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        if isinstance(line,tuple): txt,sz,b,c=line
        else: txt,sz,b,c=line,size,bold,color
        r=p.add_run(); r.text=txt; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
        p.space_after=Pt(4)
    return s
def bullets(slide, x,y,w,h, items, size=13, color=INK, gap=6):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=s.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.03)
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        pPr=p._p.get_or_add_pPr(); pPr.set('marL','228600'); pPr.set('indent','-228600')
        bu=pPr.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar',{'char':'•'}); pPr.append(bu)
        if isinstance(it,tuple):
            r=p.add_run(); r.text=it[0]; r.font.bold=True; r.font.size=Pt(size); r.font.color.rgb=color
            r=p.add_run(); r.text=it[1]; r.font.size=Pt(size); r.font.color.rgb=color
        else:
            r=p.add_run(); r.text=it; r.font.size=Pt(size); r.font.color.rgb=color
        p.space_after=Pt(gap)
    return s
def card(slide, x,y,w,h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s=slide.shapes.add_shape(shape, Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    if line: s.line.color.rgb=line; s.line.width=Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit=False
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0]=0.06
    s.text_frame.text=''
    return s
def footer(slide, n):
    tb(slide, 5.55,6.98,6.8,0.25, FOOT, 8, color=GREY, align=PP_ALIGN.RIGHT)
    tb(slide, 12.45,6.98,0.5,0.25, str(n), 8, color=GREY, align=PP_ALIGN.RIGHT)
def titled(layout_name, title, subtitle=None, n=None):
    s=prs.slides.add_slide(L[layout_name]); set_text(ph(s,0), title)
    if subtitle is not None and ph(s,13) is not None: set_text(ph(s,13), subtitle)
    if n: footer(s,n)
    return s

# ── 1 cover ──────────────────────────────────────────────────────────────
s=prs.slides.add_slide(L['Cover 1 (Lockup)'])
set_text(ph(s,0),'Data Engineering Pipeline on HPE Private Cloud AI')
set_text(ph(s,1),'Lab 18 · Curriculum 3.1 Feature Engineering & ETL/DataOps')
set_text(ph(s,23),'Airflow · Spark · EzPresto · CUDA-X (cuDF, cuML, cuVS) · MLflow · September 2026')
tb(s,5.55,6.98,6.8,0.25,FOOT,8,color=GREY,align=PP_ALIGN.RIGHT)

# ── 2 agenda ─────────────────────────────────────────────────────────────
s=prs.slides.add_slide(L['Agenda']); set_text(ph(s,0),'Agenda')
for idx,t in zip((12,13,14,15,16,17),['Business problem — who is about to cancel?','Technical problem — two systems, no training table','Technical solution — one tool per concept','Architecture on HPE Private Cloud AI','The pipeline, by row count · measured results','Deliverables and the value statement']):
    set_text(ph(s,idx),t)
footer(s,2)

# ── 3 business problem ───────────────────────────────────────────────────
s=titled('Title and Subtitle','Business problem','A subscription streaming service asks the question every subscription business asks',3)
tb(s,0.41,1.55,7.6,1.3,[('“Will this subscriber cancel in the next 30 days?”',30,True,INK)],anchor=MSO_ANCHOR.TOP)
bullets(s,0.41,2.95,7.4,3.6,[
 ('Churn is silent. ','A subscriber who has already decided to leave looks like any other row in the billing system until the cancellation lands.'),
 ('Retention only works before the decision. ','The retention team needs a ranked list of who is at risk this month, not a report of who left last month.'),
 ('The signal exists, but not where the decision is made. ','Disengagement shows up in viewing behaviour weeks before it shows up in billing. Nobody joins those two systems today.'),
 ('The ask: ','a repeatable, explainable way to rank every subscriber by 30-day cancellation risk, and to say which behaviours drive it.'),
],size=13,gap=8)
# stat cards
for i,(big,small) in enumerate([('200,000','active subscribers'),('11.9%','cancel within any 30-day window'),('23,796','cancellations in the labelled period'),('~100','viewing events per subscriber, never joined to billing')]):
    y=1.55+i*1.28; card(s,8.45,y,4.45,1.13)
    tb(s,8.6,y+0.1,4.2,0.6,big,28,True,GREEN,anchor=MSO_ANCHOR.TOP)
    tb(s,8.6,y+0.68,4.2,0.42,small,11,color=GREY)

# ── 4 technical problem ──────────────────────────────────────────────────
s=titled('Title and Subtitle','Technical problem','A model needs one row per person with the answer attached. That table does not exist.',4)
# two source cards
for i,(ttl,lines,col) in enumerate([
  ('Viewing log — the fact table',['180 daily CSV files · 849 MB','20,010,929 events, ~100 per person','who watched what, when, on which device','no outcome column'],INK),
  ('Customer records — the dimension table',['Postgres, owned by billing','200,000 rows, one per person','plan, tenure, tickets, payment, country','churned_next_30d — the only copy of the answer'],ORANGE)]):
    x=0.41+i*6.35; card(s,x,1.55,6.1,2.0)
    tb(s,x+0.2,1.65,5.8,0.4,ttl,15,True,col)
    bullets(s,x+0.2,2.05,5.8,1.5,lines,size=12,gap=3)
tb(s,0.41,3.65,12.5,0.62,'Neither source can train anything alone. The whole engineering task is to shrink the first into one row per person and attach the answer from the second, without losing the people who matter most.',13,color=GREY)
# four hard parts
for i,(ttl,body) in enumerate([
  ('Scale','20 million rows is too large for a laptop and too slow row-by-row. It needs distributed, columnar processing.'),
  ('Separation','Two systems, two owners, two formats. The label may not be copied out of the database.'),
  ('Silence','3,436 subscribers watched nothing in 30 days. 63% of them churned. A standard inner join deletes them with no error.'),
  ('Repetition','Thirty people must run the identical pipeline on their own outputs, with every run recorded and comparable.')]):
    x=0.41+i*3.17; card(s,x,4.35,3.02,2.35,fill=WHITE,line=LIGHT)
    tb(s,x+0.15,4.45,2.75,0.4,ttl,14,True,GREEN)
    tb(s,x+0.15,4.85,2.75,1.8,body,11,color=INK)

# ── 5 technical solution ─────────────────────────────────────────────────
s=titled('Title and Subtitle','Technical solution','One tool per concept, each appearing at the moment the problem makes it necessary',5)
rows=[('Airflow','Orchestration','One parameterised DAG, triggered per participant with a student number. Issues instructions, never touches a row.'),
      ('Spark on the Spark Operator','Distributed curation','180 CSV files → 180 tasks. Dedupe (shuffle), filter, write day-partitioned Parquet. Same 20,010,929 rows, 3.66× smaller.'),
      ('EzPresto','Federation','Query the billing database where it lives through a catalog. No export, no copy, read-only for everyone in the room.'),
      ('RAPIDS cuDF · cuML · cuVS','GPU acceleration, same API','Aggregate 20 M events to one row per person on the GPU (12×), train a forest (2.4×), find look-alikes (85×) — accuracy checked before the clock.'),
      ('XGBoost on GPU + MLflow','Model and registry','Held-out AUC 0.87. Every run logged with parameters, metric and model; registered as student-NN-churn, versioned.'),
      ('Validation checklist','Trust before training','Nine executable checks on the training table, including the left join that keeps the silent subscribers.')]
for i,(tool,concept,what) in enumerate(rows):
    y=1.5+i*0.87; card(s,0.41,y,12.5,0.78,fill=PANEL if i%2==0 else WHITE, line=None if i%2==0 else LIGHT)
    tb(s,0.6,y+0.08,3.3,0.62,tool,13,True,INK,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,3.95,y+0.08,2.3,0.62,concept,12,True,GREEN,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,6.3,y+0.08,6.45,0.62,what,11,color=INK,anchor=MSO_ANCHOR.MIDDLE)

# ── 6 architecture ───────────────────────────────────────────────────────
s=titled('Title Only','Architecture on HPE Private Cloud AI',None,6)
pic=s.shapes.add_picture(ARCH, Inches(0.41), Inches(1.0), width=Inches(10.5))   # 1600x900 → 10.5 x 5.91
tb(s,11.1,1.0,1.85,5.9,[('Reading the diagram',13,True,INK),
   ('Solid arrows carry rows; dashed arrows carry instructions.',11,False,GREY),
   ('Airflow sits on dashed paths only.',11,False,GREY),
   ('Storage is shared and read-only except each participant’s curated folder and MLflow experiment.',11,False,GREY),
   ('The notebook is where the two sources finally meet.',11,False,GREY),
   ('Every number on the diagram was measured on pcai1dev, seed 42.',11,False,GREY)])

# ── 7 pipeline by row count ──────────────────────────────────────────────
s=titled('Title and Subtitle','The pipeline, by row count','Four steps. Watch the row count change — that change is the data engineering.',7)
steps=[('Raw events','20,010,929','180 CSV files · 849 MB · no label',INK),
       ('Curated Parquet','20,010,929','Spark · 180 partitions · 232 MB',GREEN),
       ('The squash','196,564','cuDF groupby · 4 invented features',GREEN),
       ('Training table','200,000 × 19','left join + label · zeros for silence',ORANGE),
       ('Model','AUC 0.87','XGBoost on GPU · registered in MLflow',GREEN)]
n=len(steps); w=2.36; gap=0.2; x0=0.41
for i,(name,num,desc,col) in enumerate(steps):
    x=x0+i*(w+gap)
    c=s.shapes.add_shape(MSO_SHAPE.CHEVRON if i>0 else MSO_SHAPE.PENTAGON, Inches(x),Inches(1.6),Inches(w),Inches(0.6))
    c.fill.solid(); c.fill.fore_color.rgb=col; c.line.fill.background(); c.shadow.inherit=False
    tf=c.text_frame; tf.text=''; r=tf.paragraphs[0].add_run(); r.text=name; r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=WHITE; tf.paragraphs[0].alignment=PP_ALIGN.CENTER
    card(s,x,2.4,w,1.6)
    tb(s,x+0.1,2.5,w-0.2,0.7,num,22,True,INK,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,x+0.1,3.2,w-0.2,0.75,desc,10.5,color=GREY,align=PP_ALIGN.CENTER)
card(s,0.41,4.3,12.5,2.35,fill=WHITE,line=LIGHT)
tb(s,0.6,4.4,6.0,0.4,'The number that teaches the most',14,True,GREEN)
tb(s,0.6,4.8,12.1,1.8,[('196,564 ≠ 200,000. The 3,436 subscribers with no activity in the last 30 days have no row after the squash. 63% of them churned, against 11.9% overall.',13,False,INK),
   ('An inner join drops them silently and deletes 2,172 churners — the strongest signal in the data — while looking completely correct. The pipeline uses a LEFT join from the customer table and fills their activity with zero. Absence of a record is itself information.',12,False,GREY)])

# ── 8 measured results ───────────────────────────────────────────────────
s=titled('Title and Subtitle','Measured results','Every number below was measured on the platform, seed 42. Nothing is quoted from a datasheet.',8)
stats=[('3.66×','smaller on disk','CSV 849 MB → Parquet 232 MB, same 20,010,929 rows'),
       ('12×','faster aggregation','cuDF 0.03 s vs pandas 0.39 s, identical output'),
       ('2.4×','faster RandomForest','cuML 1.69 s vs scikit-learn 4.02 s on 4 cores · AUC gap 0.0009'),
       ('85×','faster nearest-neighbour search','cuVS 0.05 s vs scikit-learn 4.03 s · recall@10 0.9996'),
       ('0.87','held-out AUC','XGBoost on GPU · 200,000 × 19 · look-alike score alone 0.81'),
       ('0.002','importance spread of the planted noise','six country columns indistinguishable; real features differ 10×')]
for i,(big,lab,det) in enumerate(stats):
    col=i%3; row=i//3; x=0.41+col*4.2; y=1.55+row*2.45
    card(s,x,y,4.0,2.25)
    tb(s,x+0.2,y+0.15,3.6,0.8,big,36,True,GREEN,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,x+0.2,y+0.92,3.6,0.55,lab,13,True,INK)
    tb(s,x+0.2,y+1.45,3.6,0.75,det,11,color=GREY)
tb(s,0.41,6.5,12.5,0.4,'Speed-ups are always stated against a named CPU core count. Vendor headline figures are ceilings; these are the honest numbers.',11,color=GREY)

# ── 9 deliverables & value ───────────────────────────────────────────────
s=titled('Title and Subtitle','Deliverables and the value statement','What a participant leaves with, and how a measured result becomes a sentence a customer acts on',9)
for i,(ttl,body) in enumerate([
  ('Working pipeline flow','Raw events and a live database → validated training table → registered model, run end to end under the participant’s own student number.'),
  ('Transformation & validation checklist','Thirty rows across six stages; the notebook executes the table-side rows and prints PASS or FAIL for each.'),
  ('Customer-facing data-readiness page','A one-page template, filled with the participant’s own numbers, ending in a readiness verdict.'),
  ('Presales value statements','Three statements built from measured results, at least one about a data-engineering outcome rather than acceleration.')]):
    x=0.41+i*3.17; card(s,x,1.55,3.02,2.1,fill=WHITE,line=LIGHT)
    tb(s,x+0.15,1.65,2.75,0.6,ttl,13,True,GREEN)
    tb(s,x+0.15,2.25,2.75,1.35,body,11,color=INK)
card(s,0.41,3.95,12.5,2.75)
tb(s,0.6,4.05,6,0.4,'Result → meaning → value',14,True,INK)
for i,(h,b) in enumerate([('Technical result','The 20 M-row aggregation took 0.39 s on CPU and 0.03 s on the GPU, with identical output.'),
                          ('What it means','The feature-engineering loop is interactive, and the pandas code did not change.'),
                          ('Value to the customer','Your data scientists try ten feature ideas in the time it used to take to try one, with the skills they already have.')]):
    x=0.6+i*4.1; tb(s,x,4.55,3.9,0.35,h,12,True,GREEN); tb(s,x,4.9,3.9,1.7,b,12,color=INK)
tb(s,0.6,6.25,12.1,0.4,'Say only the third line to a customer. Have the first ready for when they ask “compared with what?”',11,color=GREY)

# ── 10 close ─────────────────────────────────────────────────────────────
s=prs.slides.add_slide(L['Thank You']); set_text(ph(s,0),'Thank you')
set_text(ph(s,13),'Lab 18 · Data Engineering Pipeline on HPE Private Cloud AI · guide, notebook, DAG and deliverables in the data-engineering-lab repository')
tb(s,6.12,6.75,6.8,0.25,FOOT,8,color=GREY,align=PP_ALIGN.RIGHT)

prs.save(OUT); print('saved', OUT, len(prs.slides), 'slides')
