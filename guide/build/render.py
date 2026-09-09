"""Render the lab guide content model to (a) .docx on the HPE Arial A4 template via
python-docx and (b) an HTML twin rendered to PDF with headless Chromium.
Content lives in content.py as a list of blocks:
  ('h1', text) ('h2', text) ('h3', text) ('p', text) ('lead', text)
  ('bullets', [str,...]) ('numbers', [str,...]) ('code', text) ('note', label, text)
  ('table', [header,...], [[cell,...],...], caption or None)
  ('figure', 'shots/file.png', caption) ('pagebreak',) ('cover', title, subtitle, blurb)
  ('kv', [(k, v), ...])   # two-column key/value table without header
Inline **bold** and `code` are supported in p/bullets/table cells.
"""
import os, re, sys, html, base64
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE)  # docs/
sys.path.insert(0, HERE)
from content import BLOCKS, META
GREEN=RGBColor(0x01,0xA9,0x82)

# ───────────────────────────── DOCX ─────────────────────────────
def add_runs(par, text, size=None, font=None):
    """**bold** and `code` inline markup."""
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok: continue
        if tok.startswith('**'): r=par.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith('`'): r=par.add_run(tok[1:-1]); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Consolas')
        else: r=par.add_run(tok)
        if size: r.font.size=Pt(size)
        if font and not tok.startswith('`'): r.font.name=font
    return par

PPR_ORDER=['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr','suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap','overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd','snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc','textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr','sectPr','pPrChange']
TBLPR_ORDER=['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize','tblW','jc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar','tblLook','tblCaption','tblDescription']
def insert_ordered(parent, child, order):
    """Insert child into parent at the position the OOXML schema requires."""
    tag=child.tag.split('}')[1]; idx=order.index(tag)
    for i,existing in enumerate(list(parent)):
        et=existing.tag.split('}')[1]
        if et in order and order.index(et)>idx:
            parent.insert(i, child); return
    parent.append(child)

def shade(cell, hex6):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hex6); tcPr.append(shd)

def set_cell_borders(table, hex6='B1B9BE'):
    tbl=table._tbl; tblPr=tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:space'),'0'); x.set(qn('w:color'),hex6); b.append(x)
    insert_ordered(tblPr, b, TBLPR_ORDER)

def docx_table(doc, header, rows, caption=None, kv=False):
    if caption:
        c=doc.add_paragraph(style='Caption - Table'); r=c.add_run(f'Table {caption[0]}. '); r.bold=True; c.add_run(caption[1])
    ncol=len(header) if header else len(rows[0])
    t=doc.add_table(rows=0, cols=ncol); t.alignment=WD_TABLE_ALIGNMENT.LEFT
    try: t.style=doc.styles['HPE_Table']
    except Exception: t.style=doc.styles['Table Grid']
    set_cell_borders(t)
    usable=Cm(21.0-2*1.27)
    if kv: widths=[Cm(5.0), usable-Cm(5.0)]
    else: widths=[usable/ncol]*ncol
    if header:
        cells=t.add_row().cells
        for i,h in enumerate(header):
            cells[i].width=widths[i]; shade(cells[i],'01A982'); p=cells[i].paragraphs[0]; p.style=doc.styles['Table Subhead']
            r=p.add_run(h); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        # repeat header row
        trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].width=widths[i]; p=cells[i].paragraphs[0]; p.style=doc.styles['Table text']
            add_runs(p, str(v), size=9)
            if kv and i==0:
                for r in p.runs: r.bold=True
    doc.add_paragraph()
    return t

def docx_figure(doc, path, caption, n):
    full=os.path.join(ROOT, path)
    if not os.path.exists(full):
        t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; set_cell_borders(t)
        c=t.rows[0].cells[0]; c.width=Cm(21.0-2*1.27); shade(c,'F2F4F5')
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(60); p.paragraph_format.space_after=Pt(60)
        r=p.add_run(f'Screenshot to be inserted: {os.path.basename(path)}'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x59,0x59,0x59)
        cp=doc.add_paragraph(style='Caption - Figure'); rr=cp.add_run(f'Figure {n}. '); rr.bold=True; cp.add_run(caption)
        return
    im=Image.open(full); w,h=im.size
    maxw=Cm(21.0-2*1.27); width=maxw if w>=h*0.9 else Cm(11)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    run=p.add_run(); pic=run.add_picture(full, width=width)
    # 0.75pt grey border per template guidance
    inline=run._element.find('.//'+qn('wp:inline'))
    if inline is not None:
        graphic=inline.find(qn('a:graphic')); pic_el=graphic.find('.//'+qn('pic:spPr'))
        if pic_el is not None:
            ln=OxmlElement('a:ln'); ln.set('w','9525'); sf=OxmlElement('a:solidFill'); clr=OxmlElement('a:srgbClr'); clr.set('val','B1B9BE'); sf.append(clr); ln.append(sf); pic_el.append(ln)
    c=doc.add_paragraph(style='Caption - Figure'); r=c.add_run(f'Figure {n}. '); r.bold=True; c.add_run(caption)

def new_num(doc, abstract_id, restart=True):
    """Create a fresh <w:num> for an abstractNum so each numbered list restarts at 1."""
    numbering=doc.part.numbering_part.element
    nums=numbering.findall(qn('w:num')); next_id=max(int(n.get(qn('w:numId'))) for n in nums)+1
    num=OxmlElement('w:num'); num.set(qn('w:numId'),str(next_id))
    a=OxmlElement('w:abstractNumId'); a.set(qn('w:val'),str(abstract_id)); num.append(a)
    if restart:
        lo=OxmlElement('w:lvlOverride'); lo.set(qn('w:ilvl'),'0'); so=OxmlElement('w:startOverride'); so.set(qn('w:val'),'1'); lo.append(so); num.append(lo)
    mac=numbering.find(qn('w:numIdMacAtCleanup'))
    if mac is not None: mac.addprevious(num)
    else: numbering.append(num)
    return next_id

def list_par(doc, text, num_id):
    p=doc.add_paragraph(style='List Paragraph')
    pPr=p._p.get_or_add_pPr(); numPr=OxmlElement('w:numPr')
    il=OxmlElement('w:ilvl'); il.set(qn('w:val'),'0'); ni=OxmlElement('w:numId'); ni.set(qn('w:val'),str(num_id))
    numPr.append(il); numPr.append(ni); insert_ordered(pPr, numPr, PPR_ORDER)
    add_runs(p, text); return p

def docx_code(doc, text):
    for line in text.rstrip('\n').split('\n'):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.left_indent=Cm(0.4)
        pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'F2F4F5'); insert_ordered(pPr, shd, PPR_ORDER)
        r=p.add_run(line if line else ' '); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Consolas'); r.font.size=Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def build_docx(out):
    doc=Document(os.path.join(HERE,'hpe_arial_a4_base.docx'))
    body=doc.element.body
    for p in list(body.findall(qn('w:p'))): body.remove(p)   # drop the placeholder paragraph
    fig=0; tab=0
    for b in BLOCKS:
        k=b[0]
        if k=='cover':
            _,title,subtitle,blurb=b
            for _ in range(6): doc.add_paragraph()
            p=doc.add_paragraph(style='Title'); r=p.add_run(title); r.font.size=Pt(34); r.bold=True
            p=doc.add_paragraph(style='Subtitle'); r=p.add_run(subtitle); r.font.size=Pt(16); r.font.color.rgb=GREEN
            doc.add_paragraph()
            p=doc.add_paragraph(); add_runs(p, blurb, size=11)
            for _ in range(14): doc.add_paragraph()
            p=doc.add_paragraph(); r=p.add_run(META['doc_line']); r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x59,0x59,0x59)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k=='h1':
            doc.add_paragraph(b[1], style='Heading 1')
        elif k=='h2': doc.add_paragraph(b[1], style='Heading 2')
        elif k=='h3': doc.add_paragraph(b[1], style='Heading 3')
        elif k=='p': add_runs(doc.add_paragraph(), b[1])
        elif k=='lead': add_runs(doc.add_paragraph(), b[1], size=11)
        elif k=='bullets':
            for it in b[1]: list_par(doc, it, 10)                 # numId 10 = HPE bullet
        elif k=='numbers':
            nid=new_num(doc, 19)                                   # abstractNum 19 = HPE number, restart at 1
            for it in b[1]: list_par(doc, it, nid)
        elif k=='code': docx_code(doc, b[1])
        elif k=='note':
            p=doc.add_paragraph(style='Block Text'); r=p.add_run(b[1]); r.bold=True
            p=doc.add_paragraph(style='Block Text'); add_runs(p, b[2])
        elif k=='table':
            tab+=1; docx_table(doc, b[1], b[2], (tab, b[3]) if b[3] else None)
        elif k=='kv':
            docx_table(doc, None, b[1], None, kv=True)
        elif k=='figure':
            fig+=1; docx_figure(doc, b[1], b[2], fig)
        elif k=='pagebreak':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else: raise ValueError(k)
    doc.save(out); print('docx ->', out, f'({fig} figures, {tab} tables)')

# ───────────────────────────── HTML / PDF ─────────────────────────────
def inline(text):
    t=html.escape(text)
    t=re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', t); t=re.sub(r'`([^`]+)`', r'<code>\1</code>', t)
    return t

CSS='''
@page { size: A4; margin: 18mm 12.7mm 16mm 12.7mm; }
body { font-family: Arial, "Liberation Sans", sans-serif; font-size: 10pt; color:#000; line-height:1.35; }
h1 { font-size:20pt; font-weight:bold; margin:0 0 8pt 0; page-break-before: always; padding-top:4pt; }
h1.first { page-break-before: auto; }
h2 { font-size:15pt; font-weight:bold; margin:14pt 0 6pt 0; page-break-after: avoid; }
h3 { font-size:12pt; font-weight:bold; margin:12pt 0 4pt 0; page-break-after: avoid; }
p { margin:0 0 6pt 0; }
p.lead { font-size:11pt; }
ul, ol { margin:0 0 6pt 0; padding-left:18pt; } li { margin-bottom:2pt; }
ul li::marker { color:#01A982; }
code { font-family: Consolas, "Liberation Mono", monospace; font-size:9pt; background:#F2F4F5; padding:0 2pt; }
pre { font-family: Consolas, "Liberation Mono", monospace; font-size:8.5pt; background:#F2F4F5; padding:6pt 8pt; margin:2pt 0 8pt 0; white-space:pre-wrap; border-left:3px solid #01A982; page-break-inside:avoid; }
table { border-collapse:collapse; width:100%; margin:0 0 10pt 0; font-size:9pt; page-break-inside:auto; }
th { background:#01A982; color:#fff; font-weight:bold; text-align:left; padding:4pt 6pt; border:1px solid #B1B9BE; }
td { padding:4pt 6pt; border:1px solid #B1B9BE; vertical-align:top; }
tr { page-break-inside:avoid; }
table.kv td:first-child { font-weight:bold; width:5cm; background:#F2F4F5; }
.cap { font-size:9pt; margin:3pt 0 10pt 0; } .cap b { font-weight:bold; }
.tcap { font-size:9pt; margin:8pt 0 3pt 0; }
figure { margin:8pt 0 4pt 0; page-break-inside:avoid; }
figure img { max-width:100%; border:0.75pt solid #B1B9BE; display:block; }
figure img.tall { max-width:11cm; }
.ph { border:1px dashed #B1B9BE; background:#F2F4F5; color:#595959; font-style:italic; font-size:9pt; text-align:center; padding:60pt 0; margin:8pt 0 4pt 0; page-break-inside:avoid; }
.note { border-top:1pt solid #000; border-bottom:1pt solid #000; padding:5pt 0; margin:8pt 0 10pt 0; page-break-inside:avoid; }
.note b.lab { display:block; margin-bottom:2pt; }
.cover { }
.cover .title { font-size:34pt; font-weight:bold; margin-top:60mm; line-height:1.1; }
.cover .sub { font-size:16pt; color:#01A982; margin-top:8pt; }
.cover .blurb { font-size:11pt; margin-top:18pt; max-width:150mm; }
.cover .docline { margin-top:95mm; font-size:9pt; color:#595959; }
.hdr { display:flex; justify-content:space-between; align-items:flex-start; font-size:8pt; margin-bottom:10mm; }
.hdr .logo { height:9mm; } .hdr .right { text-align:right; } .hdr .right b { font-size:11pt; display:block; }
'''
def build_html(out):
    logo_png=os.path.join(HERE,'hpe_logo.png')
    logo_tag=''
    if logo_png and os.path.exists(logo_png):
        ext=logo_png.rsplit('.',1)[1].lower(); mime={'png':'image/png','jpeg':'image/jpeg','jpg':'image/jpeg','svg':'image/svg+xml'}[ext]
        logo_tag=f'<img class="logo" src="data:{mime};base64,{base64.b64encode(open(logo_png,"rb").read()).decode()}">'
    parts=[f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(META["title"])}</title><style>{CSS}</style></head><body>']
    fig=0; tab=0; first_h1=True
    for b in BLOCKS:
        k=b[0]
        if k=='cover':
            _,title,subtitle,blurb=b
            parts.append(f'<div class="hdr">{logo_tag}<div class="right"><b>Lab guide</b>Confidential | For training purposes only</div></div>')
            parts.append(f'<div class="cover"><div class="title">{html.escape(title)}</div><div class="sub">{html.escape(subtitle)}</div><div class="blurb">{inline(blurb)}</div><div class="docline">{html.escape(META["doc_line"])}</div></div><div style="page-break-after:always"></div>')
        elif k=='h1':
            parts.append(f'<h1{" class=first" if first_h1 else ""}>{html.escape(b[1])}</h1>'); first_h1=False
        elif k=='h2': parts.append(f'<h2>{html.escape(b[1])}</h2>')
        elif k=='h3': parts.append(f'<h3>{html.escape(b[1])}</h3>')
        elif k=='p': parts.append(f'<p>{inline(b[1])}</p>')
        elif k=='lead': parts.append(f'<p class="lead">{inline(b[1])}</p>')
        elif k=='bullets': parts.append('<ul>'+''.join(f'<li>{inline(i)}</li>' for i in b[1])+'</ul>')
        elif k=='numbers': parts.append('<ol>'+''.join(f'<li>{inline(i)}</li>' for i in b[1])+'</ol>')
        elif k=='code': parts.append(f'<pre>{html.escape(b[1].rstrip())}</pre>')
        elif k=='note': parts.append(f'<div class="note"><b class="lab">{html.escape(b[1])}</b>{inline(b[2])}</div>')
        elif k=='table':
            tab+=1
            cap=f'<div class="tcap"><b>Table {tab}.</b> {html.escape(b[3])}</div>' if b[3] else ''
            parts.append(cap+'<table><thead><tr>'+''.join(f'<th>{inline(h)}</th>' for h in b[1])+'</tr></thead><tbody>'+''.join('<tr>'+''.join(f'<td>{inline(str(c))}</td>' for c in r)+'</tr>' for r in b[2])+'</tbody></table>')
        elif k=='kv':
            parts.append('<table class="kv"><tbody>'+''.join(f'<tr><td>{inline(kk)}</td><td>{inline(str(v))}</td></tr>' for kk,v in b[1])+'</tbody></table>')
        elif k=='figure':
            fig+=1; full=os.path.join(ROOT,b[1])
            if os.path.exists(full):
                im=Image.open(full); cls=' class="tall"' if im.size[0]<im.size[1]*0.9 else ''
                data=base64.b64encode(open(full,'rb').read()).decode(); ext=full.rsplit('.',1)[1].lower()
                parts.append(f'<figure><img{cls} src="data:image/{"jpeg" if ext in ("jpg","jpeg") else ext};base64,{data}"></figure><div class="cap"><b>Figure {fig}.</b> {html.escape(b[2])}</div>')
            else:
                parts.append(f'<div class="ph">Screenshot to be inserted: {html.escape(os.path.basename(b[1]))}</div><div class="cap"><b>Figure {fig}.</b> {html.escape(b[2])}</div>')
        elif k=='pagebreak': parts.append('<div style="page-break-after:always"></div>')
    parts.append('</body></html>')
    open(out,'w').write('\n'.join(parts)); print('html ->', out)

def html_to_pdf(html_path, pdf_path):
    from playwright.sync_api import sync_playwright
    exe=os.path.expanduser('~/.cache/ms-playwright/chromium-1223/chrome-linux64/chrome')
    foot='<div style="font-family:Arial,Liberation Sans,sans-serif;font-size:8pt;color:#000;width:100%;padding:0 12.7mm;text-align:right">Confidential | For training purposes only&nbsp;&nbsp;&nbsp;&nbsp;<span class="pageNumber"></span></div>'
    with sync_playwright() as p:
        b=p.chromium.launch(executable_path=exe); pg=b.new_page()
        pg.goto('file://'+os.path.abspath(html_path)); pg.wait_for_load_state('networkidle')
        pg.pdf(path=pdf_path, format='A4', print_background=True, display_header_footer=True, header_template='<div></div>', footer_template=foot,
               margin={'top':'18mm','bottom':'16mm','left':'12.7mm','right':'12.7mm'})
        b.close()
    print('pdf ->', pdf_path, os.path.getsize(pdf_path),'bytes')

if __name__=='__main__':
    out=ROOT
    base=META['file_base']
    build_docx(os.path.join(out,base+'.docx'))
    build_html(os.path.join(out,base+'.html'))
    html_to_pdf(os.path.join(out,base+'.html'), os.path.join(out,base+'.pdf'))
